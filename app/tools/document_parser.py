"""增强文档解析器 - 支持PDF、Word、Excel等"""
import os
from pathlib import Path
from typing import Dict, Any, Optional
import logging

logger = logging.getLogger(__name__)


def parse_file(file_path: str, file_type: str = None) -> Dict[str, Any]:
    """Parse various file types and extract text content.
    
    Args:
        file_path: Path to the file
        file_type: File type hint (auto-detected if not provided)
    
    Returns:
        {
            "success": bool,
            "content": str,
            "metadata": dict,
            "error": str (if failed)
        }
    """
    path = Path(file_path)
    
    if not path.exists():
        return {"success": False, "error": f"File not found: {file_path}"}
    
    # Auto-detect file type
    if not file_type:
        file_type = path.suffix.lower().lstrip(".")
    
    parsers = {
        "pdf": parse_pdf,
        "txt": parse_text,
        "md": parse_text,
        "html": parse_html,
        "py": parse_code,
        "js": parse_code,
        "ts": parse_code,
        "json": parse_json,
        "csv": parse_csv,
        "xlsx": parse_excel,
        "xls": parse_excel,
        "docx": parse_docx,
        "doc": parse_docx,
    }
    
    parser = parsers.get(file_type)
    if not parser:
        # Try as plain text
        return parse_text(file_path)
    
    try:
        return parser(file_path)
    except Exception as e:
        logger.error(f"Failed to parse {file_path}: {e}")
        return {"success": False, "error": str(e)}


def parse_pdf(file_path: str) -> Dict[str, Any]:
    """Parse PDF with multiple fallback methods."""
    content = ""
    metadata = {"pages": 0, "method": None}
    
    # Method 1: pdfplumber (best for tables)
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            pages = []
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    pages.append(text)
            content = "\n\n".join(pages)
            metadata["pages"] = len(pdf.pages)
            metadata["method"] = "pdfplumber"
    except ImportError:
        pass
    except Exception as e:
        logger.warning(f"pdfplumber failed: {e}")
    
    # Method 2: PyPDF2 (fallback)
    if not content:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            pages = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages.append(text)
            content = "\n\n".join(pages)
            metadata["pages"] = len(reader.pages)
            metadata["method"] = "PyPDF2"
        except ImportError:
            pass
        except Exception as e:
            logger.warning(f"PyPDF2 failed: {e}")
    
    # Method 3: pymupdf (another fallback)
    if not content:
        try:
            import fitz  # pymupdf
            doc = fitz.open(file_path)
            pages = []
            for page in doc:
                pages.append(page.get_text())
            content = "\n\n".join(pages)
            metadata["pages"] = len(doc)
            metadata["method"] = "pymupdf"
        except ImportError:
            pass
        except Exception as e:
            logger.warning(f"pymupdf failed: {e}")
    
    if content:
        return {
            "success": True,
            "content": content,
            "metadata": metadata,
        }
    
    return {
        "success": False,
        "error": "No PDF parser available. Install: pip install pdfplumber",
        "content": "",
        "metadata": metadata,
    }


def parse_text(file_path: str) -> Dict[str, Any]:
    """Parse plain text file."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        
        return {
            "success": True,
            "content": content,
            "metadata": {
                "lines": content.count("\n") + 1,
                "chars": len(content),
            }
        }
    except Exception as e:
        return {"success": False, "error": str(e)}


def parse_html(file_path: str) -> Dict[str, Any]:
    """Parse HTML file and extract text."""
    try:
        from bs4 import BeautifulSoup
        
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            html = f.read()
        
        soup = BeautifulSoup(html, "html.parser")
        
        # Remove script and style elements
        for script in soup(["script", "style"]):
            script.decompose()
        
        text = soup.get_text(separator="\n", strip=True)
        
        return {
            "success": True,
            "content": text,
            "metadata": {
                "title": soup.title.string if soup.title else None,
                "chars": len(text),
            }
        }
    except ImportError:
        # Fallback to regex
        import re
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            html = f.read()
        
        # Remove HTML tags
        text = re.sub(r"<[^>]+>", " ", html)
        text = re.sub(r"\s+", " ", text).strip()
        
        return {
            "success": True,
            "content": text,
            "metadata": {"chars": len(text)},
        }
    except Exception as e:
        return {"success": False, "error": str(e)}


def parse_code(file_path: str) -> Dict[str, Any]:
    """Parse code file with structure analysis."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            code = f.read()
        
        # Extract basic structure
        import re
        
        structure = []
        
        # Python
        if file_path.endswith(".py"):
            functions = re.findall(r"def (\w+)\(", code)
            classes = re.findall(r"class (\w+)", code)
            imports = re.findall(r"^(?:import|from)\s+(\S+)", code, re.MULTILINE)
            
            if imports:
                structure.append(f"Imports: {', '.join(imports[:10])}")
            if classes:
                structure.append(f"Classes: {', '.join(classes)}")
            if functions:
                structure.append(f"Functions: {', '.join(functions[:20])}")
        
        # JavaScript/TypeScript
        elif file_path.endswith((".js", ".ts")):
            functions = re.findall(r"function\s+(\w+)", code)
            classes = re.findall(r"class\s+(\w+)", code)
            exports = re.findall(r"export\s+(?:const|let|var|function|class)\s+(\w+)", code)
            
            if classes:
                structure.append(f"Classes: {', '.join(classes)}")
            if functions:
                structure.append(f"Functions: {', '.join(functions[:20])}")
            if exports:
                structure.append(f"Exports: {', '.join(exports[:10])}")
        
        header = f"# File: {Path(file_path).name}\n"
        if structure:
            header += "# " + "\n# ".join(structure) + "\n\n"
        
        return {
            "success": True,
            "content": header + code,
            "metadata": {
                "language": Path(file_path).suffix.lstrip("."),
                "lines": code.count("\n") + 1,
                "structure": structure,
            }
        }
    except Exception as e:
        return {"success": False, "error": str(e)}


def parse_json(file_path: str) -> Dict[str, Any]:
    """Parse JSON file."""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        
        # Pretty print
        content = json.dumps(data, indent=2, ensure_ascii=False)
        
        return {
            "success": True,
            "content": content,
            "metadata": {
                "type": type(data).__name__,
                "size": len(content),
            }
        }
    except Exception as e:
        return {"success": False, "error": str(e)}


def parse_csv(file_path: str) -> Dict[str, Any]:
    """Parse CSV file with preview."""
    try:
        import csv
        
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            rows = list(reader)
        
        if not rows:
            return {"success": True, "content": "(empty CSV)", "metadata": {}}
        
        # Format as table
        headers = rows[0]
        preview_rows = rows[1:21]  # First 20 rows
        
        # Calculate column widths
        widths = [len(h) for h in headers]
        for row in preview_rows:
            for i, cell in enumerate(row):
                if i < len(widths):
                    widths[i] = max(widths[i], len(str(cell)[:50]))
        
        # Build table string
        lines = []
        lines.append(" | ".join(h.ljust(widths[i]) for i, h in enumerate(headers)))
        lines.append("-+-".join("-" * w for w in widths))
        
        for row in preview_rows:
            cells = [str(row[i] if i < len(row) else "").ljust(widths[i])[:50] for i in range(len(headers))]
            lines.append(" | ".join(cells))
        
        content = "\n".join(lines)
        content += f"\n\n... Total {len(rows) - 1} rows, {len(headers)} columns"
        
        return {
            "success": True,
            "content": content,
            "metadata": {
                "rows": len(rows) - 1,
                "columns": len(headers),
                "headers": headers,
            }
        }
    except Exception as e:
        return {"success": False, "error": str(e)}


def parse_excel(file_path: str) -> Dict[str, Any]:
    """Parse Excel file."""
    try:
        import pandas as pd
        
        # Read all sheets
        excel_file = pd.ExcelFile(file_path)
        
        parts = []
        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(excel_file, sheet_name=sheet_name)
            
            parts.append(f"## Sheet: {sheet_name}")
            parts.append(f"Shape: {df.shape[0]} rows x {df.shape[1]} columns")
            parts.append(f"Columns: {', '.join(df.columns.tolist())}")
            parts.append("")
            
            # Preview
            preview = df.head(10).to_string()
            parts.append(preview)
            parts.append("")
        
        return {
            "success": True,
            "content": "\n".join(parts),
            "metadata": {
                "sheets": excel_file.sheet_names,
            }
        }
    except ImportError:
        return {
            "success": False,
            "error": "Excel parsing requires pandas. Install: pip install pandas openpyxl",
        }
    except Exception as e:
        return {"success": False, "error": str(e)}


def parse_docx(file_path: str) -> Dict[str, Any]:
    """Parse Word document."""
    try:
        from docx import Document
        
        doc = Document(file_path)
        
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)
        
        content = "\n\n".join(parts)
        
        return {
            "success": True,
            "content": content,
            "metadata": {
                "paragraphs": len(doc.paragraphs),
                "chars": len(content),
            }
        }
    except ImportError:
        return {
            "success": False,
            "error": "Word parsing requires python-docx. Install: pip install python-docx",
        }
    except Exception as e:
        return {"success": False, "error": str(e)}


# Import json for parse_json
import json
